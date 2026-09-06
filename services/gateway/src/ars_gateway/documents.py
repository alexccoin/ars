"""Learning files — PDF, DOCX, TXT, Markdown.

A document Alex uploads is his, but it is not necessarily *his words*: a contract, a
manual, a paper someone sent him. So every chunk enters memory as `USER_DATA` — trusted
as data, never as instruction. A PDF is an excellent place to hide a line addressed to an
AI assistant, and this is the door that line would come through.
"""

from __future__ import annotations

import hashlib
import json
import re
import time
from dataclasses import asdict, dataclass
from pathlib import Path

from ars_protocol import (
    Language, MemoryKind, MemoryRecord, Provenance, Sensitivity, SourceKind, TrustLevel,
    new_id,
)

SUPPORTED = {".pdf", ".txt", ".md", ".markdown", ".docx"}

CHUNK_CHARS = 1100
"""Big enough to hold a clause or a paragraph whole, small enough that a retrieval hit
points at something specific. The document tier returns the chunk verbatim, so a chunk
is also the unit Alex reads."""

CHUNK_OVERLAP = 180
"""An answer that straddles a boundary is otherwise unfindable."""


@dataclass(slots=True)
class LearnedDocument:
    id: str
    name: str
    kind: str
    chars: int
    chunks: int
    pages: int | None
    added_at_ms: int
    sha256: str


@dataclass(slots=True)
class ParsedDocument:
    name: str
    text: str
    pieces: list[str]
    pages: int | None
    language: Language
    sha256: str


class UnsupportedDocument(ValueError):
    pass


def _pdf_text(data: bytes) -> tuple[str, int]:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")  # one unreadable page must not lose the other 200
    return "\n\n".join(parts), len(reader.pages)


def _docx_text(data: bytes) -> tuple[str, None]:
    from io import BytesIO

    import docx

    d = docx.Document(BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    # Tables carry the actual content in most real-world documents — invoices, specs,
    # schedules. Dropping them loses exactly the facts people ask about.
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), None


def extract_text(name: str, data: bytes) -> tuple[str, int | None]:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise UnsupportedDocument(
            f"{suffix or 'that file type'} is not supported; "
            f"A.R.S can learn {', '.join(sorted(SUPPORTED))}"
        )
    if suffix == ".pdf":
        return _pdf_text(data)
    if suffix == ".docx":
        return _docx_text(data)
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding), None
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), None


_WS = re.compile(r"[ \t]+")
_BLANKS = re.compile(r"\n{3,}")


def normalise(text: str) -> str:
    """PDF extraction produces ragged whitespace and hyphenated line breaks. Left alone,
    'agree-\\nment' never matches a search for 'agreement'."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = _WS.sub(" ", text)
    return _BLANKS.sub("\n\n", text).strip()


def chunk(text: str, *, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split on paragraph boundaries where possible, hard-split only when forced."""
    text = normalise(text)
    if len(text) <= size:
        return [text] if text else []

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        if len(para) > size:
            if current:
                chunks.append(current)
                current = ""
            for i in range(0, len(para), size - overlap):
                chunks.append(para[i: i + size])
            continue
        if len(current) + len(para) + 2 > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            current = (tail + "\n\n" + para).strip()
        else:
            current = f"{current}\n\n{para}".strip() if current else para
    if current:
        chunks.append(current)
    return [c for c in chunks if c.strip()]


def detect_language(text: str) -> Language:
    """Cheap and good enough to tag a chunk. Romanian diacritics and a handful of very
    common function words separate the two decisively; no model needed for a label."""
    sample = text[:4000].lower()
    ro_marks = sum(sample.count(ch) for ch in "ăâîșțşţ")
    ro_words = sum(sample.count(f" {w} ") for w in
                   ("și", "este", "sunt", "pentru", "care", "din", "cu", "nu", "să", "la"))
    en_words = sum(sample.count(f" {w} ") for w in
                   ("the", "and", "is", "are", "for", "which", "from", "with", "not", "to"))
    return Language.RO if (ro_marks * 3 + ro_words) > en_words else Language.EN


CATALOGUE_PREFIX = "document:"

_LABEL = re.compile(r"^(?P<name>.+?)\s+\(\d+/\d+\)$")


def _name_from_label(label: str | None) -> str | None:
    """`store` writes chunk labels as "contract.txt (2/7)" — this reads the name back."""
    if not label:
        return None
    match = _LABEL.match(label)
    return match.group("name") if match else label


class DocumentLibrary:
    """Ingests files into the memory store and keeps a catalogue of what was learned.

    The catalogue lives in the store, not in this object. A library that empties itself on
    restart while the chunks stay searchable is the same lie in the other direction as a
    deleted document that can still be recalled: what the user is shown has to match what
    the machine actually holds. So each entry is persisted next to the records it names,
    and `rehydrate` drops any whose chunks have gone.
    """

    def __init__(self, memory) -> None:
        self.memory = memory
        self.docs: dict[str, LearnedDocument] = {}
        self._chunk_ids: dict[str, list[str]] = {}

    async def rehydrate(self) -> int:
        """Rebuild the catalogue from the store. Returns how many documents came back."""
        self.docs.clear()
        self._chunk_ids.clear()
        for key, value in await self.memory.meta_items(CATALOGUE_PREFIX):
            entry = json.loads(value)
            chunk_ids: list[str] = entry.pop("chunk_ids", [])
            alive = [i for i in chunk_ids if await self.memory.exists(i)]
            if not alive:
                # Every chunk is gone — retention swept it, or a `forget` was interrupted
                # half-way. Either way there is no document here to show.
                await self.memory.meta_delete(key)
                continue
            doc = LearnedDocument(**entry)
            self.docs[doc.id] = doc
            self._chunk_ids[doc.id] = alive
        await self._adopt_orphaned_chunks()
        return len(self.docs)

    async def _adopt_orphaned_chunks(self) -> int:
        """Rebuild entries for chunks stored before the catalogue was persisted.

        Those chunks still answer questions — they are in the index — while showing up
        nowhere in the library, which is the inconsistency this class exists to prevent.
        Everything needed is already on the records: `doc:<id>#<n>` in the uri says which
        document and which chunk, and the label carries the filename. What cannot be
        recovered is the page count and the file digest, so they come back as unknown
        rather than as a plausible-looking guess.
        """
        orphans = await self.memory.records_by_uri_prefix("doc:")
        grouped: dict[str, list] = {}
        for record in orphans:
            doc_id = (record.provenance.uri or "")[len("doc:"):].split("#")[0]
            if doc_id and doc_id not in self.docs:
                grouped.setdefault(doc_id, []).append(record)

        for doc_id, records in grouped.items():
            records.sort(key=lambda r: r.provenance.uri or "")
            name = _name_from_label(records[0].provenance.label) or "unknown document"
            doc = LearnedDocument(
                id=doc_id, name=name,
                kind=Path(name).suffix.lower().lstrip(".") or "txt",
                chars=sum(len(r.text) for r in records), chunks=len(records),
                pages=None, sha256="",
                added_at_ms=min(r.created_at_ms for r in records),
            )
            ids = [r.id for r in records]
            self.docs[doc_id] = doc
            self._chunk_ids[doc_id] = ids
            await self._persist(doc, ids)
        return len(grouped)

    async def _persist(self, doc: LearnedDocument, chunk_ids: list[str]) -> None:
        await self.memory.meta_set(
            f"{CATALOGUE_PREFIX}{doc.id}",
            json.dumps({**asdict(doc), "chunk_ids": chunk_ids}),
        )

    @staticmethod
    def parse(name: str, data: bytes) -> ParsedDocument:
        """CPU-bound half: decode, extract, chunk. Safe to run in a worker thread.

        Kept separate from `store` because the memory store's SQLite connection belongs to
        the event loop that opened it — parsing a 200-page PDF must not block that loop,
        and must not touch the database from another one either.
        """
        text, pages = extract_text(name, data)
        pieces = chunk(text)
        if not pieces:
            raise UnsupportedDocument(
                f"no readable text in {name} — if it is a scanned PDF it needs OCR first"
            )
        return ParsedDocument(name=name, text=text, pieces=pieces, pages=pages,
                              language=detect_language(text),
                              sha256=hashlib.sha256(data).hexdigest())

    async def learn(self, name: str, data: bytes) -> LearnedDocument:
        return await self.store(self.parse(name, data))

    async def store(self, parsed: ParsedDocument) -> LearnedDocument:
        name, pieces, pages = parsed.name, parsed.pieces, parsed.pages
        text, language, digest = parsed.text, parsed.language, parsed.sha256
        doc_id = new_id("doc")
        ids: list[str] = []
        for index, piece in enumerate(pieces):
            record = await self.memory.remember(MemoryRecord(
                kind=MemoryKind.DOCUMENT,
                text=piece,
                language=language,
                sensitivity=Sensitivity.PERSONAL,
                provenance=Provenance(
                    source=SourceKind.LOCAL_FILE,
                    # USER_DATA, not USER: Alex owns the file, but somebody else may have
                    # written what is inside it.
                    trust=TrustLevel.USER_DATA,
                    uri=f"doc:{doc_id}#{index}",
                    label=f"{name} ({index + 1}/{len(pieces)})",
                ),
            ))
            ids.append(record.id)

        doc = LearnedDocument(
            id=doc_id, name=name, kind=Path(name).suffix.lower().lstrip(".") or "txt",
            chars=len(text), chunks=len(pieces), pages=pages,
            added_at_ms=int(time.time() * 1000), sha256=digest,
        )
        self.docs[doc_id] = doc
        self._chunk_ids[doc_id] = ids
        await self._persist(doc, ids)
        return doc

    async def forget(self, doc_id: str) -> int:
        """Remove a document and every chunk it produced. Deleting the catalogue entry
        while leaving the chunks searchable would be a lie."""
        removed = 0
        for record_id in self._chunk_ids.pop(doc_id, []):
            removed += await self.memory.forget(record_id=record_id)
        self.docs.pop(doc_id, None)
        await self.memory.meta_delete(f"{CATALOGUE_PREFIX}{doc_id}")
        return removed

    def catalogue(self) -> list[dict]:
        return [
            {
                "id": d.id, "name": d.name, "kind": d.kind, "chars": d.chars,
                "chunks": d.chunks, "pages": d.pages, "added_at": d.added_at_ms,
            }
            for d in sorted(self.docs.values(), key=lambda x: -x.added_at_ms)
        ]
